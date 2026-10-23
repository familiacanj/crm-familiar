from fastapi import FastAPI, APIRouter, HTTPException, Request, Depends
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
import random
import string
import uuid
import httpx
from pathlib import Path
from pydantic import BaseModel
from typing import Optional, List, Literal
from datetime import datetime, timezone, timedelta

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

EMERGENT_SESSION_URL = "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data"

app = FastAPI()


@app.get("/health", include_in_schema=False)
@app.head("/health", include_in_schema=False)
async def health():
    return {"status": "ok"}


api_router = APIRouter(prefix="/api")
logger = logging.getLogger(__name__)


def now():
    return datetime.now(timezone.utc)


def new_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


NOID = {"_id": 0}

# ---------------------------------------------------------------- seed data
SEED_MEMBERS = [("Eu", "#3B825C"), ("Esposa", "#C2185B"), ("Filha", "#F57C00"), ("Mãe", "#6D4C41")]
SEED_CATEGORIES = [
    ("income_extra", "Outras rendas (extra, 13º...)", "gift"),
    ("fixed", "Aluguel / Financiamento", "home-city"),
    ("fixed", "Condomínio", "office-building"),
    ("fixed", "Água", "water"),
    ("fixed", "Luz", "flash"),
    ("fixed", "Gás", "fire"),
    ("fixed", "Internet / Celular", "wifi"),
    ("fixed", "IPTU", "file-document"),
    ("fixed", "IPVA", "car"),
    ("fixed", "Seguro (casa/carro)", "shield-check"),
    ("fixed", "Plano de saúde", "hospital-box"),
    ("fixed", "Escola / Creche", "school"),
    ("fixed", "Assinaturas (streaming etc.)", "television-play"),
    ("variable", "Cartão de crédito – Eu", "credit-card"),
    ("variable", "Cartão de crédito – Esposa", "credit-card"),
    ("variable", "Cartão de crédito – Filha", "credit-card"),
    ("variable", "Cartão de crédito – Mãe", "credit-card"),
    ("variable", "Mercado", "cart"),
    ("variable", "Combustível / Transporte", "gas-station"),
    ("variable", "Farmácia", "pill"),
    ("variable", "Lazer / Restaurantes", "party-popper"),
    ("variable", "Manutenção casa/carro", "wrench"),
    ("variable", "Outros gastos", "dots-horizontal-circle"),
]


# ---------------------------------------------------------------- models
class SessionIn(BaseModel):
    session_id: str


class FamilyCreate(BaseModel):
    name: str


class FamilyJoin(BaseModel):
    code: str


class MemberIn(BaseModel):
    name: str
    color: str = "#3B825C"
    default_income: float = 0


class CategoryIn(BaseModel):
    name: str
    kind: Literal["income_extra", "fixed", "variable"]
    icon: str = "tag"
    default_value: float = 0
    member_id: Optional[str] = None  # who this expense belongs to (e.g. credit card owner)


class CategoryUpdate(BaseModel):
    name: Optional[str] = None
    icon: Optional[str] = None
    default_value: Optional[float] = None
    member_id: Optional[str] = None  # "" clears the owner


class GoalIn(BaseModel):
    name: str
    icon: str = "airplane"
    target: float


class DepositIn(BaseModel):
    year: int
    month: int
    amount: Optional[float] = None  # None removes the deposit of that month


class EntryIn(BaseModel):
    year: int
    month: int
    ref_type: Literal["member", "category"]
    ref_id: str
    value: Optional[float] = None  # None removes override/entry
    apply_all: bool = False  # set as default monthly value (members / fixed)


# ---------------------------------------------------------------- auth
async def get_user(request: Request) -> dict:
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else None
    # File downloads (Excel export) can't always send headers, so accept ?token= on export routes only
    if not token and request.url.path.startswith("/api/export/"):
        token = request.query_params.get("token")
    if not token:
        raise HTTPException(401, "Não autenticado")
    sess = await db.user_sessions.find_one({"session_token": token}, NOID)
    if not sess:
        raise HTTPException(401, "Sessão inválida")
    exp = sess["expires_at"]
    if exp.tzinfo is None:
        exp = exp.replace(tzinfo=timezone.utc)
    if exp < now():
        raise HTTPException(401, "Sessão expirada")
    user = await db.users.find_one({"user_id": sess["user_id"]}, NOID)
    if not user:
        raise HTTPException(401, "Usuário não encontrado")
    return user


async def get_family_id(user: dict = Depends(get_user)) -> str:
    if not user.get("family_id"):
        raise HTTPException(400, "Você ainda não faz parte de uma família")
    return user["family_id"]


@api_router.post("/auth/session")
async def auth_session(body: SessionIn):
    async with httpx.AsyncClient(timeout=20) as hc:
        r = await hc.get(EMERGENT_SESSION_URL, headers={"X-Session-ID": body.session_id})
    if r.status_code != 200:
        raise HTTPException(401, "Sessão inválida")
    data = r.json()
    email = data["email"]
    existing = await db.users.find_one({"email": email}, NOID)
    if existing:
        user_id = existing["user_id"]
        await db.users.update_one({"user_id": user_id}, {"$set": {"name": data.get("name"), "picture": data.get("picture")}})
    else:
        user_id = new_id("user")
        await db.users.insert_one({"user_id": user_id, "email": email, "name": data.get("name"),
                                   "picture": data.get("picture"), "family_id": None, "created_at": now()})
    token = data["session_token"]
    await db.user_sessions.insert_one({"session_token": token, "user_id": user_id,
                                       "expires_at": now() + timedelta(days=7), "created_at": now()})
    user = await db.users.find_one({"user_id": user_id}, NOID)
    user.pop("created_at", None)
    return {"session_token": token, "user": user}


@api_router.get("/auth/me")
async def auth_me(user: dict = Depends(get_user)):
    user.pop("created_at", None)
    return user


@api_router.post("/auth/logout")
async def auth_logout(request: Request):
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        await db.user_sessions.delete_one({"session_token": auth[7:]})
    return {"ok": True}


# ---------------------------------------------------------------- family
def make_code() -> str:
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # no O/0/I/1
    return "".join(random.choice(alphabet) for _ in range(6))


CARD_PREFIX = "Cartão de crédito – "


async def seed_family(family_id: str):
    members = [
        {"member_id": new_id("mem"), "family_id": family_id, "name": n, "color": c,
         "default_income": 0.0, "order": i, "created_at": now()}
        for i, (n, c) in enumerate(SEED_MEMBERS)
    ]
    await db.members.insert_many(members)
    by_name = {m["name"]: m["member_id"] for m in members}
    await db.categories.insert_many([
        {"category_id": new_id("cat"), "family_id": family_id, "kind": k, "name": n, "icon": ic,
         "default_value": 0.0, "order": i, "created_at": now(),
         "member_id": by_name.get(n[len(CARD_PREFIX):]) if n.startswith(CARD_PREFIX) else None}
        for i, (k, n, ic) in enumerate(SEED_CATEGORIES)
    ])


async def link_card_owners():
    """Existing families: link 'Cartão de crédito – X' categories to member X."""
    async for c in db.categories.find({"member_id": {"$exists": False}, "name": {"$regex": f"^{CARD_PREFIX}"}}):
        m = await db.members.find_one({"family_id": c["family_id"], "name": c["name"][len(CARD_PREFIX):]})
        await db.categories.update_one({"_id": c["_id"]}, {"$set": {"member_id": m["member_id"] if m else None}})


@api_router.post("/family")
async def create_family(body: FamilyCreate, user: dict = Depends(get_user)):
    if user.get("family_id"):
        raise HTTPException(400, "Você já faz parte de uma família")
    code = make_code()
    while await db.families.find_one({"invite_code": code}):
        code = make_code()
    family_id = new_id("fam")
    await db.families.insert_one({"family_id": family_id, "name": body.name.strip() or "Minha Família",
                                  "invite_code": code, "created_at": now()})
    await seed_family(family_id)
    await db.users.update_one({"user_id": user["user_id"]}, {"$set": {"family_id": family_id}})
    return await family_info(family_id)


@api_router.post("/family/join")
async def join_family(body: FamilyJoin, user: dict = Depends(get_user)):
    fam = await db.families.find_one({"invite_code": body.code.strip().upper()}, NOID)
    if not fam:
        raise HTTPException(404, "Código não encontrado")
    await db.users.update_one({"user_id": user["user_id"]}, {"$set": {"family_id": fam["family_id"]}})
    return await family_info(fam["family_id"])


async def family_info(family_id: str) -> dict:
    fam = await db.families.find_one({"family_id": family_id}, {"_id": 0, "created_at": 0})
    users = await db.users.find({"family_id": family_id}, {"_id": 0, "name": 1, "email": 1, "picture": 1, "user_id": 1}).to_list(50)
    members = await db.members.find({"family_id": family_id}, {"_id": 0, "created_at": 0}).sort("order", 1).to_list(100)
    return {**fam, "users": users, "members": members}


@api_router.get("/family")
async def get_family(family_id: str = Depends(get_family_id)):
    return await family_info(family_id)


@api_router.post("/family/leave")
async def leave_family(user: dict = Depends(get_user)):
    await db.users.update_one({"user_id": user["user_id"]}, {"$set": {"family_id": None}})
    return {"ok": True}


# ---------------------------------------------------------------- members
@api_router.post("/members")
async def add_member(body: MemberIn, family_id: str = Depends(get_family_id)):
    count = await db.members.count_documents({"family_id": family_id})
    doc = {"member_id": new_id("mem"), "family_id": family_id, "name": body.name.strip(), "color": body.color,
           "default_income": max(body.default_income, 0), "order": count, "created_at": now()}
    await db.members.insert_one(doc)
    doc.pop("_id", None)
    doc.pop("created_at", None)
    return doc


@api_router.put("/members/{member_id}")
async def update_member(member_id: str, body: MemberIn, family_id: str = Depends(get_family_id)):
    res = await db.members.update_one({"member_id": member_id, "family_id": family_id},
                                      {"$set": {"name": body.name.strip(), "color": body.color,
                                                "default_income": max(body.default_income, 0)}})
    if res.matched_count == 0:
        raise HTTPException(404, "Membro não encontrado")
    return await db.members.find_one({"member_id": member_id}, {"_id": 0, "created_at": 0})


@api_router.delete("/members/{member_id}")
async def delete_member(member_id: str, family_id: str = Depends(get_family_id)):
    await db.members.delete_one({"member_id": member_id, "family_id": family_id})
    await db.entries.delete_many({"family_id": family_id, "ref_id": member_id})
    await db.categories.update_many({"family_id": family_id, "member_id": member_id}, {"$set": {"member_id": None}})
    return {"ok": True}


# ---------------------------------------------------------------- categories
@api_router.get("/categories")
async def list_categories(family_id: str = Depends(get_family_id)):
    return await db.categories.find({"family_id": family_id}, {"_id": 0, "created_at": 0}).sort("order", 1).to_list(500)


@api_router.post("/categories")
async def add_category(body: CategoryIn, family_id: str = Depends(get_family_id)):
    last = await db.categories.find({"family_id": family_id}).sort("order", -1).to_list(1)
    doc = {"category_id": new_id("cat"), "family_id": family_id, "kind": body.kind, "name": body.name.strip(),
           "icon": body.icon, "default_value": max(body.default_value, 0) if body.kind == "fixed" else 0.0,
           "member_id": body.member_id or None,
           "order": (last[0]["order"] + 1) if last else 0, "created_at": now()}
    await db.categories.insert_one(doc)
    doc.pop("_id", None)
    doc.pop("created_at", None)
    return doc


@api_router.put("/categories/{category_id}")
async def update_category(category_id: str, body: CategoryUpdate, family_id: str = Depends(get_family_id)):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    if "member_id" in upd:
        upd["member_id"] = upd["member_id"] or None
    if upd:
        res = await db.categories.update_one({"category_id": category_id, "family_id": family_id}, {"$set": upd})
        if res.matched_count == 0:
            raise HTTPException(404, "Conta não encontrada")
    return await db.categories.find_one({"category_id": category_id}, {"_id": 0, "created_at": 0})


@api_router.delete("/categories/{category_id}")
async def delete_category(category_id: str, family_id: str = Depends(get_family_id)):
    await db.categories.delete_one({"category_id": category_id, "family_id": family_id})
    await db.entries.delete_many({"family_id": family_id, "ref_id": category_id})
    return {"ok": True}


# ---------------------------------------------------------------- entries / calc
@api_router.put("/entries")
async def set_entry(body: EntryIn, family_id: str = Depends(get_family_id)):
    if not 1 <= body.month <= 12:
        raise HTTPException(400, "Mês inválido")
    if body.apply_all:
        value = max(body.value or 0, 0)
        if body.ref_type == "member":
            res = await db.members.update_one({"member_id": body.ref_id, "family_id": family_id}, {"$set": {"default_income": value}})
        else:
            res = await db.categories.update_one({"category_id": body.ref_id, "family_id": family_id, "kind": "fixed"},
                                                 {"$set": {"default_value": value}})
        if res.matched_count == 0:
            raise HTTPException(400, "Não é possível aplicar a todos os meses")
        # remove overrides from this month onward so the new monthly value applies
        await db.entries.delete_many({"family_id": family_id, "ref_id": body.ref_id,
                                      "$or": [{"year": {"$gt": body.year}},
                                              {"year": body.year, "month": {"$gte": body.month}}]})
        return {"ok": True}
    key = {"family_id": family_id, "year": body.year, "month": body.month, "ref_id": body.ref_id}
    if body.value is None:
        await db.entries.delete_one(key)
    else:
        await db.entries.update_one(key, {"$set": {"value": max(body.value, 0), "ref_type": body.ref_type,
                                                   "updated_at": now()}}, upsert=True)
    return {"ok": True}


async def load(family_id: str, year: int):
    members = await db.members.find({"family_id": family_id}, {"_id": 0, "created_at": 0}).sort("order", 1).to_list(100)
    cats = await db.categories.find({"family_id": family_id}, {"_id": 0, "created_at": 0}).sort("order", 1).to_list(500)
    entries = await db.entries.find({"family_id": family_id, "year": year}, NOID).to_list(10000)
    emap = {(e["month"], e["ref_id"]): e["value"] for e in entries}
    return members, cats, emap


def compute_month(month: int, members, cats, emap) -> dict:
    income, fixed, variable = [], [], []
    for m in members:
        ov = emap.get((month, m["member_id"]))
        income.append({"ref_type": "member", "id": m["member_id"], "name": m["name"], "icon": "account",
                       "color": m.get("color"), "kind": "member", "default_value": m["default_income"],
                       "value": ov if ov is not None else m["default_income"], "overridden": ov is not None})
    for c in cats:
        ov = emap.get((month, c["category_id"]))
        if c["kind"] == "fixed":
            val, default = (ov if ov is not None else c["default_value"]), c["default_value"]
        else:
            val, default = (ov or 0.0), 0.0
        row = {"ref_type": "category", "id": c["category_id"], "name": c["name"], "icon": c["icon"],
               "color": None, "kind": c["kind"], "default_value": default, "value": val, "overridden": ov is not None,
               "member_id": c.get("member_id")}
        {"income_extra": income, "fixed": fixed, "variable": variable}[c["kind"]].append(row)
    ti = sum(r["value"] for r in income)
    tf = sum(r["value"] for r in fixed)
    tv = sum(r["value"] for r in variable)
    te = tf + tv
    status = "" if ti == 0 and te == 0 else ("POSITIVO" if ti - te >= 0 else "NEGATIVO")
    return {"month": month, "sections": {"income": income, "fixed": fixed, "variable": variable},
            "totals": {"income": ti, "fixed": tf, "variable": tv, "expenses": te, "balance": ti - te,
                       "pct": (te / ti) if ti else None, "status": status}}


@api_router.get("/month/{year}/{month}")
async def get_month(year: int, month: int, family_id: str = Depends(get_family_id)):
    if not 1 <= month <= 12:
        raise HTTPException(400, "Mês inválido")
    members, cats, emap = await load(family_id, year)
    return {"year": year, **compute_month(month, members, cats, emap)}


@api_router.get("/year/{year}")
async def get_year(year: int, family_id: str = Depends(get_family_id)):
    members, cats, emap = await load(family_id, year)
    months, agg = [], {}
    ti = te = 0.0
    for m in range(1, 13):
        r = compute_month(m, members, cats, emap)
        t = r["totals"]
        ti += t["income"]
        te += t["expenses"]
        months.append({"month": m, "income": t["income"], "expenses": t["expenses"], "fixed": t["fixed"],
                       "variable": t["variable"], "balance": t["balance"], "status": t["status"]})
        for row in r["sections"]["fixed"] + r["sections"]["variable"]:
            a = agg.setdefault(row["id"], {"id": row["id"], "name": row["name"], "icon": row["icon"],
                                           "kind": row["kind"], "value": 0.0})
            a["value"] += row["value"]
    breakdown = sorted([a for a in agg.values() if a["value"] > 0], key=lambda a: -a["value"])
    for a in breakdown:
        a["pct"] = a["value"] / te if te else 0
    return {"year": year, "months": months,
            "totals": {"income": ti, "expenses": te, "balance": ti - te, "pct": (te / ti) if ti else None},
            "breakdown": breakdown}


@api_router.post("/month/{year}/{month}/copy-previous")
async def copy_previous(year: int, month: int, family_id: str = Depends(get_family_id)):
    """Copy last month's expenses (variable values + fixed month-only values) into empty items of this month."""
    if not 1 <= month <= 12:
        raise HTTPException(400, "Mês inválido")
    py, pm = (year - 1, 12) if month == 1 else (year, month - 1)
    cats = await db.categories.find({"family_id": family_id, "kind": {"$in": ["variable", "fixed"]}}, NOID).to_list(500)
    cat_ids = [c["category_id"] for c in cats]
    prev = await db.entries.find({"family_id": family_id, "year": py, "month": pm, "ref_id": {"$in": cat_ids}}, NOID).to_list(1000)
    current = {e["ref_id"] for e in await db.entries.find(
        {"family_id": family_id, "year": year, "month": month}, NOID).to_list(1000)}
    copied = 0
    for e in prev:
        if e["ref_id"] in current or not e.get("value"):
            continue
        await db.entries.insert_one({"family_id": family_id, "year": year, "month": month, "ref_id": e["ref_id"],
                                     "ref_type": "category", "value": e["value"], "updated_at": now()})
        copied += 1
    return {"copied": copied}


# ---------------------------------------------------------------- spending per person
@api_router.get("/people/{year}")
async def people_year(year: int, family_id: str = Depends(get_family_id)):
    members, cats, emap = await load(family_id, year)
    out = []
    for m in members:
        own = [c for c in cats if c.get("member_id") == m["member_id"] and c["kind"] in ("variable", "fixed")]
        months = []
        for mo in range(1, 13):
            total = 0.0
            for c in own:
                ov = emap.get((mo, c["category_id"]))
                total += (ov if ov is not None else c["default_value"]) if c["kind"] == "fixed" else (ov or 0.0)
            months.append(total)
        out.append({"member_id": m["member_id"], "name": m["name"], "color": m.get("color"),
                    "categories": [{"id": c["category_id"], "name": c["name"], "icon": c["icon"]} for c in own],
                    "months": months, "total": sum(months)})
    return {"year": year, "members": out}


# ---------------------------------------------------------------- savings goals
async def goal_out(g: dict) -> dict:
    deps = await db.goal_deposits.find({"goal_id": g["goal_id"]}, NOID).sort([("year", 1), ("month", 1)]).to_list(1000)
    saved = sum(d["amount"] for d in deps)
    return {"goal_id": g["goal_id"], "name": g["name"], "icon": g["icon"], "target": g["target"], "saved": saved,
            "pct": (saved / g["target"]) if g["target"] else 0,
            "deposits": [{"year": d["year"], "month": d["month"], "amount": d["amount"]} for d in deps]}


@api_router.get("/goals")
async def list_goals(family_id: str = Depends(get_family_id)):
    goals = await db.goals.find({"family_id": family_id}, NOID).sort("created_at", 1).to_list(100)
    return [await goal_out(g) for g in goals]


@api_router.post("/goals")
async def add_goal(body: GoalIn, family_id: str = Depends(get_family_id)):
    if not body.name.strip() or body.target <= 0:
        raise HTTPException(400, "Digite o nome e o valor da meta")
    g = {"goal_id": new_id("goal"), "family_id": family_id, "name": body.name.strip(), "icon": body.icon,
         "target": body.target, "created_at": now()}
    await db.goals.insert_one(g)
    return await goal_out(g)


@api_router.put("/goals/{goal_id}")
async def update_goal(goal_id: str, body: GoalIn, family_id: str = Depends(get_family_id)):
    if not body.name.strip() or body.target <= 0:
        raise HTTPException(400, "Digite o nome e o valor da meta")
    res = await db.goals.update_one({"goal_id": goal_id, "family_id": family_id},
                                    {"$set": {"name": body.name.strip(), "icon": body.icon, "target": body.target}})
    if res.matched_count == 0:
        raise HTTPException(404, "Meta não encontrada")
    return await goal_out(await db.goals.find_one({"goal_id": goal_id}, NOID))


@api_router.delete("/goals/{goal_id}")
async def delete_goal(goal_id: str, family_id: str = Depends(get_family_id)):
    await db.goals.delete_one({"goal_id": goal_id, "family_id": family_id})
    await db.goal_deposits.delete_many({"goal_id": goal_id})
    return {"ok": True}


@api_router.put("/goals/{goal_id}/deposit")
async def set_deposit(goal_id: str, body: DepositIn, family_id: str = Depends(get_family_id)):
    g = await db.goals.find_one({"goal_id": goal_id, "family_id": family_id}, NOID)
    if not g:
        raise HTTPException(404, "Meta não encontrada")
    if not 1 <= body.month <= 12:
        raise HTTPException(400, "Mês inválido")
    key = {"goal_id": goal_id, "year": body.year, "month": body.month}
    if not body.amount:
        await db.goal_deposits.delete_one(key)
    else:
        await db.goal_deposits.update_one(key, {"$set": {"amount": max(body.amount, 0)}}, upsert=True)
    return await goal_out(g)


# ---------------------------------------------------------------- export
MONTHS_SHORT = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]


async def year_table(family_id: str, year: int):
    members, cats, emap = await load(family_id, year)
    data = [compute_month(m, members, cats, emap) for m in range(1, 13)]
    blocks = []
    for key, title, total_title in [("income", "RENDA", "TOTAL RENDA"), ("fixed", "DESPESAS FIXAS", "TOTAL FIXAS"),
                                    ("variable", "DESPESAS VARIÁVEIS", "TOTAL VARIÁVEIS")]:
        rows = [(r["name"], [d["sections"][key][i]["value"] for d in data]) for i, r in enumerate(data[0]["sections"][key])]
        totals = [sum(d["sections"][key][i]["value"] for i in range(len(d["sections"][key]))) for d in data]
        blocks.append((title, rows, total_title, totals))
    summary = [
        ("TOTAL DESPESAS", [d["totals"]["expenses"] for d in data]),
        ("SALDO DO MÊS", [d["totals"]["balance"] for d in data]),
    ]
    status = [d["totals"]["status"] for d in data]
    pct = [d["totals"]["pct"] for d in data]
    return blocks, summary, status, pct


def money(v: float) -> str:
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{'-' if v < 0 else ''}R$ {s}"


@api_router.get("/export/{year}.xlsx")
async def export_xlsx(year: int, family_id: str = Depends(get_family_id)):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from fastapi.responses import StreamingResponse

    fam = await db.families.find_one({"family_id": family_id}, NOID)
    blocks, summary, status, pct = await year_table(family_id, year)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    head_fill = PatternFill("solid", fgColor="3B825C")
    total_fill = PatternFill("solid", fgColor="D1E5D7")
    bold, white = Font(bold=True), Font(bold=True, color="FFFFFF")
    fmt = '"R$" #,##0.00'
    ws.append([f"ORÇAMENTO DA FAMÍLIA – {fam['name']} – {year}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for title, rows, total_title, totals in blocks:
        ws.append([title, *MONTHS_SHORT, "Total ano"])
        for c in ws[ws.max_row]:
            c.fill, c.font = head_fill, white
        for name, vals in rows:
            ws.append([name, *vals, sum(vals)])
        ws.append([total_title, *totals, sum(totals)])
        for c in ws[ws.max_row]:
            c.fill, c.font = total_fill, bold
        ws.append([])
    for name, vals in summary:
        ws.append([name, *vals, sum(vals)])
        for c in ws[ws.max_row]:
            c.font = bold
    ws.append(["SITUAÇÃO", *status])
    ws.append(["% da renda gasta", *[p if p is not None else "" for p in pct]])
    for row in ws.iter_rows(min_row=3):
        for c in row[1:]:
            if isinstance(c.value, (int, float)):
                c.number_format = "0%" if row[0].value == "% da renda gasta" else fmt
            c.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 13

    goals = await db.goals.find({"family_id": family_id}, NOID).to_list(100)
    if goals:
        wg = wb.create_sheet("Metas")
        wg.append(["Meta", "Valor da meta", "Guardado", "% atingido"])
        for c in wg[1]:
            c.fill, c.font = head_fill, white
        for g in goals:
            o = await goal_out(g)
            wg.append([o["name"], o["target"], o["saved"], o["pct"]])
            wg.cell(wg.max_row, 2).number_format = fmt
            wg.cell(wg.max_row, 3).number_format = fmt
            wg.cell(wg.max_row, 4).number_format = "0%"
        wg.column_dimensions["A"].width = 30

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="orcamento_familia_{year}.xlsx"'})


@api_router.get("/export/{year}/html")
async def export_html(year: int, family_id: str = Depends(get_family_id)):
    from html import escape
    fam = await db.families.find_one({"family_id": family_id}, NOID)
    blocks, summary, status, pct = await year_table(family_id, year)

    def cells(vals, total=True):
        tds = "".join(f"<td>{money(v) if v else '–'}</td>" for v in vals)
        return tds + (f"<td><b>{money(sum(vals))}</b></td>" if total else "")

    head = "<tr><th></th>" + "".join(f"<th>{m}</th>" for m in MONTHS_SHORT) + "<th>Total</th></tr>"
    body = ""
    for title, rows, total_title, totals in blocks:
        body += f"<tr class='sec'><td colspan='14'>{title}</td></tr>"
        body += "".join(f"<tr><td>{escape(n)}</td>{cells(v)}</tr>" for n, v in rows)
        body += f"<tr class='tot'><td>{total_title}</td>{cells(totals)}</tr>"
    for n, v in summary:
        body += f"<tr class='tot'><td>{n}</td>{cells(v)}</tr>"
    body += "<tr><td>SITUAÇÃO</td>" + "".join(
        f"<td class='{'pos' if s == 'POSITIVO' else 'neg' if s else ''}'>{s or '–'}</td>" for s in status) + "<td></td></tr>"
    body += "<tr><td>% da renda gasta</td>" + "".join(
        f"<td>{round(p * 100)}%</td>" if p is not None else "<td>–</td>" for p in pct) + "<td></td></tr>"
    html = f"""<html><head><meta charset='utf-8'><style>
    @page {{ size: A4 landscape; margin: 12mm; }}
    body {{ font-family: Helvetica, Arial, sans-serif; color: #1A1A19; }}
    h1 {{ color: #3B825C; font-size: 20px; margin: 0 0 8px; }}
    table {{ border-collapse: collapse; width: 100%; font-size: 8.5px; }}
    th {{ background: #3B825C; color: #fff; padding: 4px; }}
    td {{ border-bottom: 1px solid #E3E2DD; padding: 3px 4px; text-align: right; white-space: nowrap; }}
    td:first-child {{ text-align: left; }}
    tr.sec td {{ background: #E8F3EB; font-weight: bold; text-align: left; }}
    tr.tot td {{ font-weight: bold; background: #F0EFEB; }}
    .pos {{ color: #2E7D32; font-weight: bold; }} .neg {{ color: #D32F2F; font-weight: bold; }}
    </style></head><body><h1>Orçamento da Família – {escape(fam['name'])} – {year}</h1>
    <table>{head}{body}</table></body></html>"""
    return {"html": html}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')


@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.users.create_index("user_id", unique=True)
    await db.user_sessions.create_index("session_token", unique=True)
    await db.user_sessions.create_index("user_id")
    await db.user_sessions.create_index("expires_at", expireAfterSeconds=0)
    await db.families.create_index("invite_code", unique=True)
    await db.entries.create_index([("family_id", 1), ("year", 1), ("month", 1), ("ref_id", 1)], unique=True)
    await db.goal_deposits.create_index([("goal_id", 1), ("year", 1), ("month", 1)], unique=True)
    await link_card_owners()


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
